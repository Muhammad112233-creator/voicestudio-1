#!/usr/bin/env python3
"""Local-only inference bridge for VoiceStudio.

This script never downloads model files and contains no network client. It accepts one
JSON request on stdin, writes progress/result JSON lines to stdout, and uses only local
executables, Python packages, and model paths configured by the user.
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import wave
from pathlib import Path
from typing import Any

_XTTS_CACHE: dict[tuple[str, str, str], Any] = {}


def emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False), flush=True)


def progress(value: int, message: str) -> None:
    emit({"type": "progress", "progress": max(0, min(100, int(value))), "message": message})


def package_available(*names: str) -> bool:
    return any(importlib.util.find_spec(name) is not None for name in names)


def probe() -> dict[str, Any]:
    return {
        "ok": True,
        "python": sys.version.split()[0],
        "platform": sys.platform,
        "packages": {
            "faster_whisper": package_available("faster_whisper"),
            "piper": package_available("piper") or shutil.which("piper") is not None,
            "coqui_tts": package_available("TTS"),
            "argostranslate": package_available("argostranslate"),
        },
        "tools": {"ffmpeg": shutil.which("ffmpeg") is not None},
        "offline": True,
    }


def clean_name(value: str, fallback: str = "voicestudio-output") -> str:
    name = Path(value or fallback).name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return (name or fallback)[:150]


def output_path(request: dict[str, Any], default_name: str) -> Path:
    root = Path(request.get("outputDir") or Path.cwd()).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    name = clean_name(request.get("output_name", default_name), default_name)
    return root / name


def require_file(value: str, what: str) -> Path:
    if not value:
        raise RuntimeError(f"Select a local {what} in Settings or import it first.")
    path = Path(value).expanduser()
    if not path.is_file():
        raise RuntimeError(f"The selected {what} does not exist: {path}")
    return path


def require_dir(value: str, what: str) -> Path:
    if not value:
        raise RuntimeError(f"Select a local {what} folder in Settings first.")
    path = Path(value).expanduser()
    if not path.is_dir():
        raise RuntimeError(f"The selected {what} folder does not exist: {path}")
    return path


def ffmpeg_executable(settings: dict[str, Any]) -> str:
    selected = str(settings.get("ffmpegPath", "") or "").strip()
    if selected:
        if Path(selected).is_file():
            return selected
        # A command can be specified by name (e.g. "ffmpeg") as well as by full path.
        found = shutil.which(selected)
        if found:
            return found
        raise RuntimeError("The configured ffmpeg executable could not be found. Check Settings.")
    found = shutil.which("ffmpeg")
    if not found:
        raise RuntimeError("ffmpeg is required for this media job. Install it locally and select it in Settings.")
    return found


def get_whisper(settings: dict[str, Any]):
    model_dir = require_dir(str(settings.get("whisperModelPath", "")), "faster-whisper model")
    if not package_available("faster_whisper"):
        raise RuntimeError("faster-whisper is not installed in the selected Python environment. Install it locally, then check setup again.")
    from faster_whisper import WhisperModel  # type: ignore

    device_choice = str(settings.get("device", "auto") or "auto").lower()
    device = "cuda" if device_choice == "cuda" else ("cpu" if device_choice == "mps" else device_choice)
    compute_type = "float16" if device == "cuda" else "int8" if device == "cpu" else "default"
    try:
        # A directory path + local_files_only prevents model lookup/download from a hub.
        return WhisperModel(str(model_dir), device=device, compute_type=compute_type, local_files_only=True)
    except TypeError:
        # Very old builds may not expose local_files_only. Do not use those builds for
        # inference: fail closed instead of risking a model resolver network request.
        raise RuntimeError("This faster-whisper build does not support local_files_only. Update the local package before using it offline.")
    except Exception as exc:
        raise RuntimeError(f"Could not load the local Whisper model: {exc}") from exc


def srt_timestamp(seconds: float) -> str:
    total_ms = max(0, int(round(seconds * 1000)))
    hours, rem = divmod(total_ms, 3_600_000)
    minutes, rem = divmod(rem, 60_000)
    secs, millis = divmod(rem, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def write_srt(segments: list[dict[str, Any]], path: Path) -> None:
    lines: list[str] = []
    for index, segment in enumerate(segments, 1):
        lines.extend([
            str(index),
            f"{srt_timestamp(float(segment.get('start', 0)))} --> {srt_timestamp(float(segment.get('end', 0)))}",
            str(segment.get("text", "")).strip(),
            "",
        ])
    path.write_text("\n".join(lines), encoding="utf-8")


def transcribe(input_file: str, settings: dict[str, Any]) -> dict[str, Any]:
    source = require_file(input_file, "audio or video file")
    model = get_whisper(settings)
    progress(9, "Loaded local speech-recognition model")
    try:
        segment_iter, info = model.transcribe(
            str(source), beam_size=5, vad_filter=True, word_timestamps=False,
        )
    except Exception as exc:
        raise RuntimeError(f"Could not decode this media file locally: {exc}. Check that ffmpeg/PyAV can read it.") from exc

    segments: list[dict[str, Any]] = []
    text_parts: list[str] = []
    for index, item in enumerate(segment_iter):
        text = str(item.text).strip()
        if not text:
            continue
        segment = {"start": float(item.start), "end": float(item.end), "text": text}
        segments.append(segment)
        text_parts.append(text)
        if index and index % 8 == 0:
            progress(min(84, 12 + index // 2), f"Recognized {len(segments)} speech segments")
    progress(88, "Writing transcript and subtitles")
    out = output_path({"outputDir": str(Path(input_file).parent)}, Path(input_file).stem + ".srt")
    write_srt(segments, out)
    return {
        "text": " ".join(text_parts).strip(),
        "language": getattr(info, "language", "") or "",
        "segments": segments,
        "output_path": str(out),
        "output_name": out.name,
        "message": f"Transcription complete · {len(segments)} segments · local subtitles saved.",
    }


def installed_translation(text: str, from_code: str, to_code: str) -> str:
    if from_code == "auto" or not from_code:
        raise RuntimeError("Choose the source language explicitly; local translation does not guess it.")
    if from_code == to_code:
        return text
    if not package_available("argostranslate"):
        raise RuntimeError("Argos Translate is not installed. Install it and the required language pack locally; VoiceStudio will not download either.")
    try:
        import argostranslate.translate  # type: ignore
        return argostranslate.translate.translate(text, from_code, to_code)
    except Exception as exc:
        raise RuntimeError(f"No installed local Argos translation package supports {from_code} → {to_code}: {exc}") from exc


def translate_request(request: dict[str, Any]) -> dict[str, Any]:
    from_code = str(request.get("from_language") or "auto")
    to_code = str(request.get("to_language") or "en")
    segments = request.get("segments") or []
    if segments:
        translated_segments = []
        for index, seg in enumerate(segments):
            translated_segments.append({
                **seg,
                "text": installed_translation(str(seg.get("text", "")), from_code, to_code),
            })
            if index % 4 == 0:
                progress(min(88, 10 + index * 2), f"Translated {index + 1}/{len(segments)} segments locally")
        text = " ".join(seg["text"] for seg in translated_segments)
        out = output_path(request, "translation.srt")
        write_srt(translated_segments, out)
        return {"text": text, "segments": translated_segments, "output_path": str(out), "output_name": out.name, "message": "Translation complete using installed local language packs."}
    source_text = str(request.get("text", ""))
    if not source_text.strip():
        raise RuntimeError("There is no transcript text to translate.")
    progress(20, "Using installed local translation package")
    translated = installed_translation(source_text, from_code, to_code)
    return {"text": translated, "message": "Translation complete using an installed local language pack."}


def xtts_engine(settings: dict[str, Any], device: str):
    model_dir = require_dir(str(settings.get("xttsModelPath", "")), "XTTS model")
    config_file = require_file(str(settings.get("xttsConfigPath", "")), "XTTS config.json")
    if not package_available("TTS"):
        raise RuntimeError("Coqui TTS is not installed in the selected Python environment. Install it locally and point to a local XTTS checkpoint/config.")
    key = (str(model_dir), str(config_file), device)
    if key not in _XTTS_CACHE:
        try:
            # Use Coqui's manual model API with an explicit local checkpoint directory.
            # This avoids the model-name resolver used by TTS("model_name"), which can download weights.
            from TTS.tts.configs.xtts_config import XttsConfig  # type: ignore
            from TTS.tts.models.xtts import Xtts  # type: ignore
            config = XttsConfig()
            config.load_json(str(config_file))
            model = Xtts.init_from_config(config)
            model.load_checkpoint(config, checkpoint_dir=str(model_dir), use_deepspeed=False)
            if device == "cuda":
                model.cuda()
            else:
                model.cpu()
            import torch  # type: ignore
            import torchaudio  # type: ignore
            _XTTS_CACHE[key] = (model, torch, torchaudio)
        except Exception as exc:
            raise RuntimeError(f"Could not load the local XTTS model/checkpoint: {exc}") from exc
    return _XTTS_CACHE[key]


def piper_executable(settings: dict[str, Any]) -> str:
    selected = str(settings.get("piperPath", "") or "").strip()
    if selected:
        if Path(selected).is_file():
            return selected
        found = shutil.which(selected)
        if found:
            return found
        raise RuntimeError("The configured Piper executable was not found. Check the path in Settings.")
    found = shutil.which("piper")
    if found:
        return found
    raise RuntimeError("Piper is not configured. Set a Piper executable and local .onnx voice model in Settings.")


def synthesize(text: str, output: Path, settings: dict[str, Any], voice: dict[str, Any] | None, language: str) -> None:
    if not text.strip():
        raise RuntimeError("There is no text to synthesize.")
    voice = voice or {}
    reference = str(voice.get("reference_path") or voice.get("referencePath") or "")
    if reference:
        reference_file = require_file(reference, "consented reference sample")
        device = str(settings.get("device", "auto") or "auto").lower()
        model, torch, torchaudio = xtts_engine(settings, device)
        try:
            gpt_cond_latent, speaker_embedding = model.get_conditioning_latents(audio_path=[str(reference_file)])
            language_code = (language if language and language != "auto" else "en")
            if language_code == "zh":
                language_code = "zh-cn"
            generated = model.inference(text, language_code, gpt_cond_latent, speaker_embedding)
            waveform = torch.tensor(generated["wav"]).unsqueeze(0).cpu()
            torchaudio.save(str(output), waveform, 24000)
        except Exception as exc:
            raise RuntimeError(f"Local XTTS synthesis failed: {exc}") from exc
        return

    model = require_file(str(settings.get("piperModelPath", "")), "Piper .onnx voice model")
    executable = piper_executable(settings)
    args = [executable, "--model", str(model), "--output_file", str(output)]
    try:
        result = subprocess.run(args, input=text.encode("utf-8"), stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    except OSError as exc:
        raise RuntimeError(f"Could not run local Piper: {exc}") from exc
    if result.returncode != 0:
        message = result.stderr.decode("utf-8", "replace")[-1200:].strip()
        raise RuntimeError(f"Piper returned code {result.returncode}: {message or 'check the local voice model'}")
    if not output.is_file() or output.stat().st_size == 0:
        raise RuntimeError("Piper did not produce an audio file. Check the local model and executable.")


def run_tts(request: dict[str, Any]) -> dict[str, Any]:
    text = str(request.get("text", ""))
    language = str(request.get("language") or "en")
    voice = request.get("voice") or {}
    out = output_path(request, "speech.wav")
    if out.suffix.lower() != ".wav":
        out = out.with_suffix(".wav")
    progress(12, "Loading selected local speech engine")
    synthesize(text, out, request.get("settings") or {}, voice, language)
    progress(92, "Audio rendered")
    return {"output_path": str(out), "output_name": out.name, "message": "Speech sample generated locally. Review and disclose synthetic audio where appropriate."}


def text_chunks(text: str, limit: int = 1100) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs or [text.strip()]:
        while len(paragraph) > limit:
            cut = paragraph.rfind(" ", 0, limit)
            if cut < limit // 2:
                cut = limit
            part, paragraph = paragraph[:cut].strip(), paragraph[cut:].strip()
            if current:
                chunks.append(current)
                current = ""
            if part:
                chunks.append(part)
        if len(current) + len(paragraph) + 2 > limit and current:
            chunks.append(current)
            current = paragraph
        else:
            current = f"{current}\n\n{paragraph}".strip()
    if current:
        chunks.append(current)
    return chunks


def merge_wavs(files: list[Path], output: Path) -> None:
    if not files:
        raise RuntimeError("There was no audio to combine.")
    try:
        with wave.open(str(files[0]), "rb") as first:
            params = first.getparams()
            frames = [first.readframes(first.getnframes())]
        for path in files[1:]:
            with wave.open(str(path), "rb") as source:
                if source.getparams()[:4] != params[:4]:
                    raise RuntimeError("The selected speech engine produced incompatible WAV formats. Export chapters separately or use a consistent model.")
                frames.append(source.readframes(source.getnframes()))
        with wave.open(str(output), "wb") as merged:
            merged.setparams(params)
            for chunk in frames:
                merged.writeframes(chunk)
    except wave.Error as exc:
        raise RuntimeError(f"Could not combine the generated WAV files: {exc}") from exc


def run_audiobook(request: dict[str, Any]) -> dict[str, Any]:
    text = str(request.get("text", "")).strip()
    if not text:
        raise RuntimeError("Add text before generating an audiobook.")
    settings = request.get("settings") or {}
    voice = request.get("voice") or {}
    language = str(request.get("language") or "en")
    out = output_path(request, "voicestudio-audiobook.wav")
    if out.suffix.lower() != ".wav":
        out = out.with_suffix(".wav")
    chunks = text_chunks(text)
    if len(chunks) > 300:
        raise RuntimeError("This text is too long for one job. Split it into chapters and render separately.")
    with tempfile.TemporaryDirectory(prefix="voicestudio-book-") as temp_dir:
        parts: list[Path] = []
        for index, chunk in enumerate(chunks):
            progress(8 + int(75 * index / max(1, len(chunks))), f"Rendering chapter chunk {index + 1}/{len(chunks)}")
            part = Path(temp_dir) / f"part-{index:04d}.wav"
            synthesize(chunk, part, settings, voice, language)
            parts.append(part)
        merge_wavs(parts, out)
    progress(94, "Combined audiobook audio")
    return {"output_path": str(out), "output_name": out.name, "message": f"Audiobook audio generated locally · {len(chunks)} chunk(s)."}


def wav_duration(path: Path) -> float:
    try:
        with wave.open(str(path), "rb") as audio:
            return audio.getnframes() / max(1, audio.getframerate())
    except (wave.Error, OSError) as exc:
        raise RuntimeError(f"The local TTS engine did not create a readable WAV: {exc}") from exc


def atempo_filter(rate: float) -> str:
    # ffmpeg atempo accepts 0.5..2.0 per stage; chain stages for extreme changes.
    value = max(0.1, min(10.0, rate))
    chain: list[str] = []
    while value > 2.0:
        chain.append("atempo=2.0")
        value /= 2.0
    while value < 0.5:
        chain.append("atempo=0.5")
        value /= 0.5
    chain.append(f"atempo={value:.6f}")
    return ",".join(chain)


def run_dub(request: dict[str, Any]) -> dict[str, Any]:
    source = require_file(str(request.get("input_path", "")), "video or audio file")
    settings = request.get("settings") or {}
    voice = request.get("voice") or {}
    if not (voice.get("reference_path") or voice.get("referencePath")):
        raise RuntimeError("Dubbing requires a voice profile with a consented reference sample. A descriptive profile is not a trained voice.")
    progress(3, "Transcribing source with local faster-whisper")
    transcription = transcribe(str(source), settings)
    segments = transcription["segments"]
    if not segments:
        raise RuntimeError("No speech was recognized in the source file.")
    if len(segments) > 300:
        raise RuntimeError("This media contains too many speech segments for one dubbing job. Split it into shorter clips.")
    target = str(request.get("target_language") or "en")
    translate = bool(request.get("translate", True))
    source_language = str(transcription.get("language") or "")
    if translate and source_language and source_language != target:
        progress(12, f"Translating speech {source_language} → {target} with installed local packs")
        for index, segment in enumerate(segments):
            segment["text"] = installed_translation(segment["text"], source_language, target)
            if index % 5 == 0:
                progress(min(28, 14 + index), f"Translated {index + 1}/{len(segments)} lines")
    elif translate and not source_language:
        raise RuntimeError("The local recognizer did not identify the source language; dubbing translation needs a known source language.")

    ffmpeg = ffmpeg_executable(settings)
    output = output_path(request, source.stem + f"-{target}-dubbed.mp4")
    if output.suffix.lower() not in {".mp4", ".mkv", ".mov"}:
        output = output.with_suffix(".mp4")
    with tempfile.TemporaryDirectory(prefix="voicestudio-dub-") as temp_dir:
        parts: list[tuple[Path, float, float]] = []
        for index, segment in enumerate(segments):
            start = max(0.0, float(segment.get("start", 0)))
            end = max(start + 0.2, float(segment.get("end", start + 0.2)))
            part = Path(temp_dir) / f"line-{index:04d}.wav"
            synthesize(str(segment.get("text", "")), part, settings, voice, target)
            duration = max(0.1, wav_duration(part))
            parts.append((part, start, end))
            progress(30 + int(52 * (index + 1) / len(segments)), f"Generating dubbed line {index + 1}/{len(segments)}")

        args = [ffmpeg, "-y", "-i", str(source)]
        for part, _start, _end in parts:
            args.extend(["-i", str(part)])
        filters: list[str] = []
        labels: list[str] = []
        for index, (part, start, end) in enumerate(parts):
            target_duration = max(0.2, end - start)
            actual = max(0.1, wav_duration(part))
            rate = actual / target_duration
            delay_ms = max(0, int(round(start * 1000)))
            label = f"v{index}"
            filters.append(
                f"[{index + 1}:a]aresample=48000,{atempo_filter(rate)},apad,atrim=duration={target_duration:.3f},adelay={delay_ms}|{delay_ms}[{label}]"
            )
            labels.append(f"[{label}]")
        filters.append(f"{''.join(labels)}amix=inputs={len(labels)}:duration=longest:normalize=0,alimiter=limit=0.97[aout]")
        args.extend(["-filter_complex", ";".join(filters), "-map", "0:v?", "-map", "[aout]", "-c:v", "copy", "-c:a", "aac", "-b:a", "192k", str(output)])
        result = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
        if result.returncode != 0 or not output.is_file():
            detail = result.stderr.decode("utf-8", "replace")[-1800:].strip()
            raise RuntimeError(f"ffmpeg could not mux the local dub: {detail or 'check media codecs and ffmpeg'}")
    progress(96, "Wrote dubbed media locally")
    return {
        "output_path": str(output), "output_name": output.name,
        "message": f"Local dub complete · {len(segments)} timed lines. Original source audio is muted in this first version.",
        "segments": segments,
    }


def extract_document(request: dict[str, Any]) -> dict[str, Any]:
    source = require_file(str(request.get("input_path", "")), "document")
    suffix = source.suffix.lower()
    progress(15, "Reading local document")
    if suffix in {".txt", ".md"}:
        text = source.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".docx":
        if not package_available("docx"):
            raise RuntimeError("DOCX import needs python-docx in the selected Python environment. Install it locally; the app will not fetch packages.")
        from docx import Document  # type: ignore
        document = Document(str(source))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts)
    elif suffix == ".pdf":
        text = ""
        if package_available("pypdf"):
            from pypdf import PdfReader  # type: ignore
            reader = PdfReader(str(source))
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        elif package_available("fitz"):
            import fitz  # type: ignore
            document = fitz.open(str(source))
            text = "\n\n".join(page.get_text("text").strip() for page in document)
            document.close()
        else:
            raise RuntimeError("PDF import needs pypdf or PyMuPDF in the selected Python environment. Scanned/image-only pages also need a separate local OCR tool.")
    else:
        raise RuntimeError("Supported documents are TXT, Markdown, DOCX and selectable-text PDF.")
    text = text.strip()
    if not text:
        raise RuntimeError("No selectable text was found. Scanned PDF pages need local OCR before import.")
    progress(92, "Document text extracted")
    return {"text": text, "message": f"Extracted {len(text)} characters locally from {source.name}."}


def execute(request: dict[str, Any]) -> dict[str, Any]:
    action = str(request.get("action", ""))
    if action == "probe":
        return probe()
    if action == "transcribe":
        return transcribe(str(request.get("input_path", "")), request.get("settings") or {})
    if action == "translate":
        return translate_request(request)
    if action == "extract_document":
        return extract_document(request)
    if action == "tts":
        return run_tts(request)
    if action == "audiobook":
        return run_audiobook(request)
    if action == "dub":
        return run_dub(request)
    raise RuntimeError(f"Unsupported local job type: {action or '(empty)'}")


def main() -> int:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--probe", action="store_true")
    args, _ = parser.parse_known_args()
    if args.probe:
        print(json.dumps(probe(), ensure_ascii=False))
        return 0
    try:
        request = json.load(sys.stdin)
        result = execute(request)
        emit({"type": "result", "ok": True, **result})
        return 0
    except Exception as exc:
        emit({"type": "result", "ok": False, "error": str(exc)})
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
